"""Importe un releve de valeurs liquidatives au format tableur.

Le format attendu est celui du fichier de reference `data/raw/vl/*.xlsx` :
une feuille « Serie », une ligne par couple (support, repere), colonnes ISIN,
Nom, Repere demande, Date retenue, Valeur, Devise, Source, Remarque.

Trois regles portent tout l'outil.

**La convention se lit dans la remarque, ligne par ligne.** Un meme fichier
melange les deux : VL quotidienne non ajustee pour les fonds, cours ajuste des
dividendes pour les actions et les ETF. Ecrire un fichier sous une convention
unique melangerait deux definitions dans une meme serie.

**C'est la date retenue qui est stockee, jamais la date visee.** Le repere « il
y a 24 mois » des actions vaut au 01/10/2024, douze jours apres la cible, et
c'est une ouverture : la date reelle et le moment sont conserves tels quels.

**Un releve ponctuel ne remplace jamais une valeur deja presente.** Il comble
un trou, sinon rien. Une serie quotidienne validee ne se laisse pas reecrire
par une transcription de tableur ; le recoupement devient un echantillon de
controle plutot qu'un ecrasement silencieux.

    python tools/importer_vl.py                 # simulation, rien n'est ecrit
    python tools/importer_vl.py --ecrire
    python tools/importer_vl.py chemin.xlsx --ecrire --refaire
"""
import datetime
import glob
import hashlib
import os
import re
import sqlite3
import sys

import openpyxl

RACINE = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
DOSSIER = os.path.join(RACINE, "data/raw/vl")
COTATIONS = os.path.join(RACINE, "data/cotations.db")
REFERENCE = os.path.join(RACINE, "data/reference.db")
SCHEMA = os.path.join(RACINE, "db/schema_cotations.sql")
RE_EXTRACTION = re.compile(r"extraites? le \w+ (\d{2}/\d{2}/\d{4})", re.I)


def convention(remarque):
    """('brute'|'ajustee', 'cloture'|'ouverture') tels que la source les declare."""
    texte = str(remarque or "").lower()
    return ("brute" if "non ajust" in texte else "ajustee",
            "ouverture" if "ouverture" in texte else "cloture")


def extrait_le(classeur):
    """Date d'extraction declaree dans la feuille « Methode », si elle y figure."""
    if "Méthode" not in classeur.sheetnames:
        return None
    for (cellule,) in classeur["Méthode"].iter_rows(values_only=True):
        trouve = RE_EXTRACTION.search(str(cellule or ""))
        if trouve:
            jour, mois, annee = trouve.group(1).split("/")
            return f"{annee}-{mois}-{jour}"
    return None


def lire(chemin):
    """(lignes, date d'extraction). Une ligne est un dictionnaire deja normalise."""
    classeur = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    if "Série" not in classeur.sheetnames:
        raise SystemExit(f"{chemin} : feuille « Série » absente.")
    lignes = []
    for brute in list(classeur["Série"].iter_rows(values_only=True))[1:]:
        if not brute or not brute[0]:
            continue
        isin, nom, repere, date, valeur, devise, source, remarque = (list(brute) + [None] * 8)[:8]
        base, moment = convention(remarque)
        lignes.append({"isin": str(isin).strip(), "nom": nom, "repere": repere,
                       "date": date.date().isoformat() if hasattr(date, "date") else None,
                       "valeur": valeur, "devise": devise, "source": source,
                       "base": base, "moment": moment})
    date_extraction = extrait_le(classeur)
    classeur.close()
    return lignes, date_extraction


def empreinte(chemin):
    return hashlib.sha256(open(chemin, "rb").read()).hexdigest()


def importer(chemin, cx, univers, ecrire):
    lignes, date_extraction = lire(chemin)
    retenues, rejets, controles = [], [], []
    for ligne in lignes:
        if ligne["isin"] not in univers:
            rejets.append((ligne["isin"], str(ligne["repere"]), "ISIN absent du référentiel",
                           str(ligne["source"] or "")[:200]))
        elif ligne["date"] is None or ligne["valeur"] in (None, ""):
            rejets.append((ligne["isin"], str(ligne["repere"]), "valeur absente à la source",
                           str(ligne["source"] or "")[:200]))
        else:
            existe = cx.execute(
                "SELECT valeur FROM valeur_liquidative WHERE isin=? AND date=? AND base=?",
                (ligne["isin"], ligne["date"], ligne["base"])).fetchone()
            if existe:
                ecart = (float(ligne["valeur"]) / existe[0] - 1) * 100 if existe[0] else 0.0
                controles.append((ligne["isin"], ligne["base"], ligne["date"], ecart))
                rejets.append((ligne["isin"], ligne["date"], "déjà présent — valeur conservée",
                               f"relevé {float(ligne['valeur']):.4f} contre {existe[0]:.4f} "
                               f"en base, écart {ecart:+.2f} %"))
            else:
                retenues.append(ligne)

    lot = None
    if ecrire:
        lot = cx.execute(
            "INSERT INTO import_cotation (fichier, empreinte, extrait_le, importe_le,"
            " lignes_lues, retenues, rejetees) VALUES (?,?,?,?,?,?,?)",
            (os.path.relpath(chemin, RACINE), empreinte(chemin), date_extraction,
             datetime.datetime.now().isoformat(timespec="seconds"),
             len(lignes), len(retenues), len(rejets))).lastrowid
        cx.executemany(
            "INSERT INTO valeur_liquidative (isin, date, base, valeur, devise, moment, source,"
            " import_id) VALUES (?,?,?,?,?,?,?,?) ON CONFLICT (isin, date, base) DO NOTHING",
            [(l["isin"], l["date"], l["base"], float(l["valeur"]), l["devise"], l["moment"],
              str(l["source"] or "relevé tableur")[:120], lot) for l in retenues])
        cx.executemany("INSERT INTO rejet_cotation (import_id, isin, date_visee, motif, detail)"
                       " VALUES (?,?,?,?,?)", [(lot, *r) for r in rejets])
        cx.commit()
    return lignes, retenues, rejets, controles, date_extraction, lot


def main():
    arguments = sys.argv[1:]
    ecrire = "--ecrire" in arguments
    refaire = "--refaire" in arguments
    chemins = [a for a in arguments if not a.startswith("--")] or sorted(
        glob.glob(os.path.join(DOSSIER, "*.xlsx")))
    if not chemins:
        raise SystemExit(f"Aucun fichier dans {os.path.relpath(DOSSIER, RACINE)}.")

    reference = sqlite3.connect(REFERENCE)
    univers = {r[0]: r[1] for r in reference.execute("SELECT isin, nom FROM instrument")}
    reference.close()

    cx = sqlite3.connect(COTATIONS)
    cx.executescript(open(SCHEMA, encoding="utf-8").read())

    for chemin in chemins:
        print(f"\n=== {os.path.relpath(chemin, RACINE)} ===")
        deja = cx.execute("SELECT id, importe_le FROM import_cotation WHERE empreinte = ?",
                          (empreinte(chemin),)).fetchone()
        if deja and not refaire:
            print(f"Déjà importé le {deja[1]} (lot {deja[0]}). --refaire pour rejouer.")
            continue
        lignes, retenues, rejets, controles, extraction, lot = importer(
            chemin, cx, univers, ecrire)

        print(f"Extrait le {extraction or 'date non déclarée'} · {len(lignes)} lignes lues")
        conventions = {}
        for l in retenues:
            conventions[l["base"]] = conventions.get(l["base"], 0) + 1
        print(f"  retenues  {len(retenues):>4}   " +
              ", ".join(f"{n} en {b}" for b, n in sorted(conventions.items())))
        ouvertures = sum(1 for l in retenues if l["moment"] == "ouverture")
        if ouvertures:
            print(f"            dont {ouvertures} relevés d'ouverture, non des clôtures")
        motifs = {}
        for _, _, motif, _ in rejets:
            motifs[motif] = motifs.get(motif, 0) + 1
        print(f"  écartées  {len(rejets):>4}")
        for motif, n in sorted(motifs.items(), key=lambda x: -x[1]):
            print(f"      {n:>4}  {motif}")

        if controles:
            ecarts = sorted((abs(e), i, b, d, e) for i, b, d, e in controles)
            au_dela = [x for x in ecarts if x[0] > 0.5]
            milieu = ecarts[len(ecarts) // 2][0]
            print(f"\n  Recoupement sur {len(controles)} points déjà en base :")
            print(f"      écart absolu médian {milieu:.3f} %, "
                  f"{len(au_dela)} au-delà de 0,5 %")
            for a, isin, base, date, ecart in reversed(ecarts[-5:]):
                print(f"      {isin}  {base:<8} {date}  {ecart:+.2f} %  {univers[isin][:30]}")

        nouveaux = sorted({l["isin"] for l in retenues})
        if nouveaux:
            print(f"\n  {len(nouveaux)} supports alimentés :")
            for isin in nouveaux:
                n = sum(1 for l in retenues if l["isin"] == isin)
                print(f"      {isin}  {n} point(s)  {univers[isin][:44]}")
        print(f"\n  {'écrit, lot ' + str(lot) if ecrire else 'SIMULATION — rien écrit'}")

    if ecrire:
        total = cx.execute("SELECT COUNT(*) FROM valeur_liquidative").fetchone()[0]
        supports = cx.execute("SELECT COUNT(DISTINCT isin) FROM valeur_liquidative").fetchone()[0]
        print(f"\nBase : {total} valeurs sur {supports} supports")
    cx.close()


if __name__ == "__main__":
    main()
